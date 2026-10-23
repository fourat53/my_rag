import io
import os

from typing import List, Tuple
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import GoogleGenerativeAIEmbeddings, ChatGoogleGenerativeAI
from langchain_pinecone import PineconeVectorStore
from langchain_core.prompts import ChatPromptTemplate
from pypdf import PdfReader

MAX_CHUNKS_FOR_DELETE = 10_000
PINECONE_INDEX_DIMENSION = 1024
gemini_api_key = os.getenv("GEMINI_API_KEY")

embeddings = GoogleGenerativeAIEmbeddings(
    model="models/gemini-embedding-001",
    google_api_key=gemini_api_key,
    output_dimensionality=PINECONE_INDEX_DIMENSION,
)

vector_store = PineconeVectorStore(index_name="chat-example", embedding=embeddings)

llm = ChatGoogleGenerativeAI(
    model="gemini-2.5-flash-lite",
    temperature=0.2,
    google_api_key=gemini_api_key,
)


def process_text(text: str) -> str:
    lines = text.split("\n")
    cleaned_lines = []
    seen = set()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if len(stripped.split()) < 4:
            continue
        if stripped not in seen:
            seen.add(stripped)
            cleaned_lines.append(stripped)
    cleaned_text = " ".join(cleaned_lines)
    return cleaned_text


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages: List[str] = []
    for page in reader.pages:
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(page_text)
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    document = DocxDocument(io.BytesIO(data))
    parts: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cells).strip()
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def extract_text_from_file(filename: str, data: bytes) -> str:
    """Ingest step: extract plain text from an uploaded document."""
    extension = os.path.splitext(filename)[1].lower()
    if extension == ".pdf":
        text = _extract_pdf(data)
    elif extension == ".docx":
        text = _extract_docx(data)
    elif extension in {".txt", ".md", ".csv"}:
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: {extension or 'unknown'}")
    return text


def add_text_to_db(index_id: str, raw_text: str) -> Tuple[List[str], int]:
    cleaned_text = process_text(raw_text)
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
    chunks = splitter.split_text(cleaned_text)

    docs = [
        Document(page_content=chunk, metadata={"index_id": index_id})
        for chunk in chunks
    ]
    if docs:
        vector_store.add_documents(docs)
    return chunks, len(chunks)


def delete_text_from_db(index_id: str) -> Tuple[List[str], int]:
    docs = vector_store.similarity_search(
        " ", k=MAX_CHUNKS_FOR_DELETE, filter={"index_id": index_id}
    )
    deleted_chunks = [doc.page_content for doc in docs]
    vector_store.delete(filter={"index_id": index_id})
    return deleted_chunks, len(deleted_chunks)


def similarity_search_chunks(
    query: str, k: int = 5, index_ids: List[str] | None = None
) -> List[Tuple[str, float]]:
    filter_dict = {"index_id": {"$in": index_ids}} if index_ids else None
    results = vector_store.similarity_search_with_score(query, k=k, filter=filter_dict)
    return [(doc.page_content, float(score)) for doc, score in results]


def query_db(question: str, index_ids: List[str] | None = None) -> str:
    filter_dict = {"index_id": {"$in": index_ids}} if index_ids else None
    docs = vector_store.similarity_search(question, k=5, filter=filter_dict)
    context = "\n\n".join([doc.page_content for doc in docs])

    prompt = ChatPromptTemplate.from_messages(
        [
            (
                "system",
                (
                    "You are a precise and reliable assistant. "
                    "Your task is to answer questions using only the information provided in the context, "
                    "which consists of multiple retrieved text chunks. "
                    "Do not introduce external knowledge, assumptions, or additional details. "
                    "Synthesize the relevant information across chunks into a clear, natural, human-like answer. "
                    "Avoid lists, bullet points, markdown formatting, symbols, or emphasis markers. "
                    "Write clean sentences, one after the other, as a short paragraph. "
                    "If the context does not contain the answer or any part of it, respond exactly with: "
                    '"I do not have enough information to answer the question."'
                    "\n\nContext:\n{context}"
                ),
            ),
            (
                "user",
                (
                    "This is an example of how to answer a question using the context.\n"
                    "Question:\n"
                    "What technologies power the RAG microservice?\n\n"
                    "Context:\n"
                    "Chunk 1: The RAG microservice runs on an Nginx server and is implemented using Python and FastAPI.\n"
                    "Chunk 2: Vector storage is handled by ChromaDB in production and FAISS during development.\n"
                    "Chunk 3: The system uses OpenAI's text-embedding-ada-002 model for embedding generation and "
                    "integrates retrieval and generation through LangChain."
                ),
            ),
            (
                "assistant",
                (
                    "The RAG microservice is built with Python and FastAPI, running on an Nginx server. "
                    "For vector storage, it uses ChromaDB in production and FAISS during development. "
                    "The system generates embeddings using OpenAI's text-embedding-ada-002 model and "
                    "integrates retrieval and generation through the LangChain library."
                ),
            ),
            (
                "user",
                "Question:\n{question}",
            ),
        ]
    )

    chain = prompt | llm
    res = chain.invoke({"context": context, "question": question})
    return res.content
