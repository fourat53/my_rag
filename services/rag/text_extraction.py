import csv
import io
import os

from collections.abc import Callable, Iterable
from docx import Document as DocxDocument
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape


def decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def extract_plain_text(data: bytes) -> str:
    return decode_text(data)


def extract_csv(data: bytes) -> str:
    rows: list[str] = []
    for row in csv.reader(io.StringIO(decode_text(data))):
        cells = [" ".join(cell.split()) for cell in row]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page in reader.pages:
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(page_text)
    return "\n\n".join(pages)


def extract_docx(data: bytes) -> str:
    document = DocxDocument(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cells).strip()
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def collect_shape_lines(shapes: Iterable[BaseShape]) -> list[str]:
    lines: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.extend(collect_shape_lines(shape.shapes))
            continue
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = " ".join(paragraph.text.split())
                if text:
                    lines.append(text)
        if shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
    return lines


def extract_pptx(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for slide in presentation.slides:
        slide_lines = collect_shape_lines(slide.shapes)
        if slide.has_notes_slide:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
            if notes:
                slide_lines.append(f"Notes: {notes}")
        if slide_lines:
            parts.append("\n".join(slide_lines))
    return "\n\n".join(parts)


EXTRACTORS: dict[str, Callable[[bytes], str]] = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".csv": extract_csv,
    ".txt": extract_plain_text,
    ".md": extract_plain_text,
}

SUPPORTED_EXTENSIONS = frozenset(EXTRACTORS)


def extract_text(filename: str, data: bytes) -> str:
    extension = os.path.splitext(filename)[1].lower()
    extractor = EXTRACTORS.get(extension)
    if extractor is None:
        raise ValueError(f"Unsupported file type: {extension or 'unknown'}")
    return extractor(data)
